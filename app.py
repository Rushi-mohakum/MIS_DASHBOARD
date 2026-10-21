"""MIS Intelligence Studio: business MIS dashboards plus data-quality and AI tools."""
from __future__ import annotations

import os
import pickle
import tempfile
import uuid
from datetime import datetime
from io import BytesIO

import pandas as pd
from flask import Flask, flash, jsonify, redirect, render_template, request, send_file, session, url_for
from werkzeug.utils import secure_filename

from ai_service import answer_dataset_question, generate_insights
from analytics import build_chart_series, build_quality_report, clean_dataframe
from parser import build_excel_report, compute_dashboard_data, load_dataframe, prepare_business_dataframe
from reports import build_pdf_report
from production_analytics import load_excel_workbook, detect_workbook_type, build_production_dashboard_data

app = Flask(__name__)
app.secret_key = os.environ.get("MIS_SECRET_KEY", "change-this-secret-in-render")
app.config["MAX_CONTENT_LENGTH"] = int(os.environ.get("MAX_UPLOAD_MB", "50")) * 1024 * 1024

ALLOWED_EXT = {"xlsx", "xls", "csv", "pdf", "docx"}
CACHE_DIR = os.path.join(tempfile.gettempdir(), "mis_intelligence_studio")
os.makedirs(CACHE_DIR, exist_ok=True)


def _cache_path(token: str) -> str:
    return os.path.join(CACHE_DIR, f"{token}.pkl")


def _save_new_cache(df: pd.DataFrame, tax_rate: float, filename: str, warnings: list[str], dataset_type: str = "business", workbook_sheets: dict | None = None, detection: dict | None = None) -> dict:
    # Keep one active dataset of each type so uploading production does not erase
    # the user's last sales workbook, and vice versa.
    cache_key = "production_dataset_token" if dataset_type == "production" else "business_dataset_token"
    old_token = session.get(cache_key)
    if old_token:
        try:
            os.unlink(_cache_path(old_token))
        except OSError:
            pass

    token = uuid.uuid4().hex
    payload = {
        "raw_df": df.copy(),
        "df": df.copy(),
        "tax_rate": tax_rate,
        "filename": filename,
        "warnings": warnings,
        "cleaning_log": [],
        "uploaded_at": datetime.now().isoformat(timespec="minutes"),
        "last_insights": None,
        "dataset_type": dataset_type,
        "workbook_sheets": workbook_sheets or {},
        "detection": detection or {},
    }
    with open(_cache_path(token), "wb") as file_obj:
        pickle.dump(payload, file_obj)
    session[cache_key] = token
    session["dataset_token"] = token
    return payload


def _load_cache(dataset_type: str | None = None) -> dict | None:
    if dataset_type == "production":
        token = session.get("production_dataset_token") or session.get("dataset_token")
    elif dataset_type == "business":
        token = session.get("business_dataset_token") or session.get("dataset_token")
    else:
        token = session.get("dataset_token")
    if not token:
        return None
    path = _cache_path(token)
    if not os.path.exists(path):
        return None
    try:
        with open(path, "rb") as file_obj:
            payload = pickle.load(file_obj)
        stored_type = payload.get("dataset_type", "business")
        if dataset_type and stored_type != dataset_type:
            return None
        return payload
    except (OSError, pickle.PickleError, EOFError):
        return None


def _write_cache(payload: dict, dataset_type: str | None = None) -> None:
    if dataset_type == "production":
        token = session.get("production_dataset_token") or session.get("dataset_token")
    elif dataset_type == "business":
        token = session.get("business_dataset_token") or session.get("dataset_token")
    else:
        token = session.get("dataset_token")
    if not token:
        raise RuntimeError("No active dataset session.")
    with open(_cache_path(token), "wb") as file_obj:
        pickle.dump(payload, file_obj)


def allowed_file(filename: str) -> bool:
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXT


def compact_currency(value) -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return "—"
    sign = "-" if number < 0 else ""
    number = abs(number)
    if number >= 10_000_000:
        return f"{sign}₹{number / 10_000_000:.2f} Cr"
    if number >= 100_000:
        return f"{sign}₹{number / 100_000:.2f} L"
    if number >= 1_000:
        return f"{sign}₹{number / 1_000:.1f} K"
    return f"{sign}₹{number:,.0f}"


def compact_number(value) -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return "—"
    if abs(number) >= 10_000_000:
        return f"{number / 10_000_000:.2f} Cr"
    if abs(number) >= 100_000:
        return f"{number / 100_000:.2f} L"
    if abs(number) >= 1_000:
        return f"{number / 1_000:.1f} K"
    return f"{number:,.0f}"


app.jinja_env.filters["compact_currency"] = compact_currency
app.jinja_env.filters["compact_number"] = compact_number


def _dashboard_context(cached: dict) -> dict:
    source_df = cached["df"]
    data = compute_dashboard_data(source_df, tax_rate=cached["tax_rate"])
    quality = build_quality_report(source_df)
    prepared_df, _ = prepare_business_dataframe(source_df)

    date_label = "All available periods"
    if "date" in prepared_df.columns:
        parsed_dates = pd.to_datetime(prepared_df["date"], errors="coerce").dropna()
        if not parsed_dates.empty:
            start, end = parsed_dates.min(), parsed_dates.max()
            date_label = start.strftime("%d %b %Y") if start == end else f"{start.strftime('%d %b %Y')} – {end.strftime('%d %b %Y')}"

    preview_df = source_df.head(10).copy()
    preview_df = preview_df.where(pd.notna(preview_df), "")

    warnings = list(dict.fromkeys(cached.get("warnings", []) + data.get("schema_warnings", [])))
    return {
        "data": data,
        "quality": quality,
        "filename": cached["filename"],
        "warnings": warnings,
        "tax_rate": round(cached["tax_rate"] * 100),
        "row_count": len(source_df),
        "columns": [str(c) for c in source_df.columns],
        "numeric_columns": quality["numeric_columns"],
        "preview_columns": [str(c) for c in preview_df.columns],
        "preview_rows": preview_df.to_dict(orient="records"),
        "cleaning_log": cached.get("cleaning_log", []),
        "uploaded_at": cached.get("uploaded_at"),
        "date_label": date_label,
        "has_api_key": bool(os.environ.get("GEMINI_API_KEY")),
        "gemini_model": os.environ.get("GEMINI_MODEL", "gemini-3.6-flash"),
        "last_insights": cached.get("last_insights"),
    }


@app.errorhandler(413)
def file_too_large(_error):
    flash(f"The uploaded file is larger than the {app.config['MAX_CONTENT_LENGTH'] // (1024 * 1024)} MB limit.", "error")
    return redirect(url_for("upload_page"))


@app.route("/", methods=["GET"])
def upload_page():
    business_cache = _load_cache("business")
    production_cache = _load_cache("production")
    return render_template(
        "upload.html",
        has_api_key=bool(os.environ.get("GEMINI_API_KEY")),
        has_dataset=business_cache is not None or production_cache is not None,
        has_business_dataset=business_cache is not None,
        has_production_dataset=production_cache is not None,
    )


@app.route("/dashboard", methods=["GET", "POST"])
def dashboard():
    if request.method == "GET":
        cached = _load_cache("business")
        if not cached:
            flash("Upload a sales or business file to create the Business MIS Dashboard.", "info")
            return redirect(url_for("upload_page"))
        return render_template("dashboard.html", **_dashboard_context(cached))

    file = request.files.get("file")
    if not file or file.filename == "":
        flash("Please choose a file to upload.", "error")
        return redirect(url_for("upload_page"))
    if not allowed_file(file.filename):
        flash("Unsupported file type. Upload .xlsx, .xls, .csv, .pdf, or .docx.", "error")
        return redirect(url_for("upload_page"))

    try:
        tax_rate = float(request.form.get("tax_rate") or 25) / 100
        tax_rate = min(max(tax_rate, 0), 1)
    except ValueError:
        tax_rate = 0.25

    original_name = secure_filename(file.filename) or f"dataset-{uuid.uuid4().hex[:8]}"
    ext = original_name.rsplit(".", 1)[1].lower()
    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=f".{ext}") as tmp:
            file.save(tmp.name)
            tmp_path = tmp.name

        if ext in ("xlsx", "xls"):
            workbook_sheets, raw_sheets = load_excel_workbook(tmp_path)
            detection = detect_workbook_type(workbook_sheets, raw_sheets)
            if detection.get("type") == "production":
                main_df = next(
                    (frame for name, frame in workbook_sheets.items() if str(name).strip().lower() == "production"),
                    None,
                )
                if main_df is None:
                    main_df = max(workbook_sheets.values(), key=lambda item: item.shape[0] * max(item.shape[1], 1), default=None)
                if main_df is None or main_df.empty:
                    flash("The workbook looks like a production file, but no usable production table was found.", "error")
                    return redirect(url_for("upload_page"))
                warnings = [
                    f"Automatically detected a production workbook with {len(workbook_sheets)} non-empty worksheets.",
                    "Production, quality, workforce, planning, mold, energy and incident metrics are mapped from the available sheets.",
                ]
                _save_new_cache(
                    main_df, tax_rate, original_name, warnings,
                    dataset_type="production", workbook_sheets=workbook_sheets, detection=detection,
                )
                flash("Production workbook detected. Opening Production Analytics Dashboard.", "success")
                return redirect(url_for("production_dashboard"))
            df, warnings = load_dataframe(tmp_path, ext)
            if detection.get("confidence") == "low":
                warnings.append("File type was not conclusive, so it was opened in Business MIS. Check the detected columns.")
        else:
            df, warnings = load_dataframe(tmp_path, ext)
            if df is not None and not df.empty:
                table_detection = detect_workbook_type({"Uploaded Table": df})
                if table_detection.get("type") == "production" and table_detection.get("confidence") != "low":
                    production_warnings = [
                        "Automatically detected a production table from its column headers.",
                        "Only metrics available in this single table can be calculated; multi-sheet modules may be unavailable.",
                    ]
                    _save_new_cache(
                        df, tax_rate, original_name, production_warnings,
                        dataset_type="production", workbook_sheets={"Production": df}, detection=table_detection,
                    )
                    flash("Production data detected from the table. Opening Production Analytics Dashboard.", "success")
                    return redirect(url_for("production_dashboard"))
    except Exception as exc:
        flash(f"Could not read that file: {exc}", "error")
        return redirect(url_for("upload_page"))
    finally:
        if tmp_path and os.path.exists(tmp_path):
            os.unlink(tmp_path)

    if df is None or df.empty:
        flash("No usable table was found. Try an Excel or CSV file containing column headings and data rows.", "error")
        return redirect(url_for("upload_page"))

    _save_new_cache(df, tax_rate, original_name, warnings, dataset_type="business")
    flash("Sales/business dataset detected. Opening Business MIS Dashboard.", "success")
    return redirect(url_for("dashboard"))


@app.route("/production-dashboard", methods=["GET"])
def production_dashboard():
    cached = _load_cache("production")
    if not cached:
        return render_template(
            "production_dashboard.html",
            has_production=False,
            production=None,
            filename=None,
            uploaded_at=None,
            has_api_key=bool(os.environ.get("GEMINI_API_KEY")),
        )
    production = build_production_dashboard_data(cached.get("workbook_sheets", {}), cached.get("filename", ""))
    return render_template(
        "production_dashboard.html",
        has_production=bool(production.get("ok")),
        production=production,
        filename=cached.get("filename"),
        uploaded_at=cached.get("uploaded_at"),
        has_api_key=bool(os.environ.get("GEMINI_API_KEY")),
    )


@app.route("/clean", methods=["POST"])
def clean_data():
    cached = _load_cache("business")
    if not cached:
        flash("Your dataset session expired. Upload the file again.", "error")
        return redirect(url_for("upload_page"))
    if cached.get("dataset_type") == "production":
        flash("Data cleaning is currently available for Business MIS datasets. Your production workbook remains unchanged.", "info")
        return redirect(url_for("production_dashboard"))

    action = request.form.get("action", "")
    try:
        if action == "reset":
            cached["df"] = cached["raw_df"].copy()
            message = "Reset all cleaning changes and restored the original uploaded data."
        else:
            cached["df"], message = clean_dataframe(cached["df"], action)
        cached.setdefault("cleaning_log", []).append({
            "time": datetime.now().strftime("%H:%M"),
            "action": action.replace("_", " ").title(),
            "message": message,
        })
        cached["last_insights"] = None
        _write_cache(cached, "business")
        flash(message, "success")
    except ValueError as exc:
        flash(str(exc), "error")
    return redirect(url_for("dashboard") + "#cleaning")


@app.route("/chart-data", methods=["POST"])
def chart_data():
    cached = _load_cache("business")
    if not cached:
        return jsonify({"ok": False, "error": "Dataset session expired."}), 400
    if cached.get("dataset_type") == "production":
        return jsonify({"ok": False, "error": "Chart Builder is available for Business MIS datasets. Use the Production Analytics Dashboard charts for this workbook."}), 400
    payload = request.get_json(silent=True) or request.form
    try:
        chart = build_chart_series(
            cached["df"],
            str(payload.get("x", "")),
            str(payload.get("y", "")),
            str(payload.get("aggregation", "sum")),
        )
        return jsonify({"ok": True, **chart})
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400


@app.route("/insights", methods=["POST"])
def insights():
    cached = _load_cache("business")
    if not cached:
        return jsonify({"ok": False, "error": "Dataset session expired. Upload the file again."}), 400
    if cached.get("dataset_type") == "production":
        return jsonify({"ok": False, "error": "Business AI Insights is not available for production workbooks yet. Use the Production Analytics Dashboard metrics."}), 400
    data = compute_dashboard_data(cached["df"], tax_rate=cached["tax_rate"])
    quality = build_quality_report(cached["df"])
    text, source = generate_insights(data, quality, cached["tax_rate"])
    cached["last_insights"] = text
    _write_cache(cached, "business")
    return jsonify({"ok": True, "text": text, "source": source})


@app.route("/chat", methods=["POST"])
def dataset_chat():
    cached = _load_cache("business")
    if not cached:
        return jsonify({"ok": False, "error": "Dataset session expired. Upload the file again."}), 400
    if cached.get("dataset_type") == "production":
        return jsonify({"ok": False, "error": "Business AI Chat is for sales/business datasets. Use the production charts and tables for this workbook."}), 400
    body = request.get_json(silent=True) or {}
    question = str(body.get("question", "")).strip()
    if not question:
        return jsonify({"ok": False, "error": "Enter a question about the dataset."}), 400
    if len(question) > 1000:
        return jsonify({"ok": False, "error": "Question is too long."}), 400

    data = compute_dashboard_data(cached["df"], tax_rate=cached["tax_rate"])
    quality = build_quality_report(cached["df"])
    answer, source = answer_dataset_question(cached["df"], question, data, quality)
    return jsonify({"ok": source != "error", "answer": answer, "source": source})


@app.route("/download/xlsx", methods=["GET"])
def download_xlsx():
    cached = _load_cache("business")
    if not cached:
        flash("Your dataset session expired. Upload the file again.", "error")
        return redirect(url_for("upload_page"))
    data = compute_dashboard_data(cached["df"], tax_rate=cached["tax_rate"])
    buffer = build_excel_report(cached["df"], data, cached["tax_rate"])
    return send_file(buffer, as_attachment=True, download_name="MIS_Intelligence_Report.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/download/cleaned.csv", methods=["GET"])
def download_cleaned_csv():
    cached = _load_cache("business")
    if not cached:
        flash("Your dataset session expired. Upload the file again.", "error")
        return redirect(url_for("upload_page"))
    if cached.get("dataset_type") == "production":
        flash("This export is for Business MIS datasets. Use Export Production Summary on the production dashboard.", "info")
        return redirect(url_for("production_dashboard"))
    buffer = BytesIO(cached["df"].to_csv(index=False).encode("utf-8-sig"))
    return send_file(buffer, as_attachment=True, download_name="MIS_Cleaned_Data.csv", mimetype="text/csv")


@app.route("/download/cleaned.xlsx", methods=["GET"])
def download_cleaned_xlsx():
    cached = _load_cache("business")
    if not cached:
        flash("Your dataset session expired. Upload the file again.", "error")
        return redirect(url_for("upload_page"))
    if cached.get("dataset_type") == "production":
        flash("This export is for Business MIS datasets. Use Export Production Summary on the production dashboard.", "info")
        return redirect(url_for("production_dashboard"))
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        cached["df"].to_excel(writer, sheet_name="Cleaned Data", index=False)
    buffer.seek(0)
    return send_file(buffer, as_attachment=True, download_name="MIS_Cleaned_Data.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/download/pdf", methods=["GET"])
def download_pdf():
    cached = _load_cache("business")
    if not cached:
        flash("Your dataset session expired. Upload the file again.", "error")
        return redirect(url_for("upload_page"))
    if cached.get("dataset_type") == "production":
        flash("This export is for Business MIS datasets. Use Export Production Summary on the production dashboard.", "info")
        return redirect(url_for("production_dashboard"))
    data = compute_dashboard_data(cached["df"], tax_rate=cached["tax_rate"])
    quality = build_quality_report(cached["df"])
    insights_text = cached.get("last_insights")
    if not insights_text:
        insights_text, _ = generate_insights(data, quality, cached["tax_rate"])
    buffer = build_pdf_report(data, quality, cached["filename"], insights_text)
    return send_file(buffer, as_attachment=True, download_name="MIS_Intelligence_Report.pdf", mimetype="application/pdf")


@app.route("/download/production-summary.xlsx", methods=["GET"])
def download_production_summary():
    cached = _load_cache("production")
    if not cached or cached.get("dataset_type") != "production":
        flash("Upload a production workbook before exporting a production summary.", "info")
        return redirect(url_for("upload_page"))

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    production = build_production_dashboard_data(cached.get("workbook_sheets", {}), cached.get("filename", ""))
    if not production.get("ok"):
        flash(production.get("error", "The production workbook could not be summarized."), "error")
        return redirect(url_for("production_dashboard"))
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Production Summary"
    summary.append(["Production Analytics Dashboard"])
    summary.append(["Source file", cached.get("filename", "")])
    summary.append(["Reporting period", production.get("period_label", "Not available")])
    summary.append(["Generated at", datetime.now().strftime("%Y-%m-%d %H:%M")])
    summary.append([])
    summary.append(["KPI", "Value"])
    for key, label in [
        ("total_production", "Total Production"),
        ("ok_production", "OK Production"),
        ("rejected_qty", "Rejected Quantity"),
        ("production_target", "Production Target"),
        ("achievement_pct", "Target Achievement %"),
        ("rejection_pct", "Rejection %"),
        ("oee_pct", "Average OEE %"),
        ("workforce_count", "Operators & Inspectors Recorded"),
        ("machine_count", "Machines Recorded"),
        ("plant_count", "Plants Recorded"),
        ("downtime_minutes", "Downtime Minutes"),
        ("material_used", "Actual Material Used"),
        ("material_required", "Material Required"),
        ("quality_ppm", "Calculated PPM"),
        ("incident_count", "Incident Records"),
        ("open_incidents", "Open Incidents"),
    ]:
        summary.append([label, production.get("kpi", {}).get(key, 0)])

    def add_table(sheet_name, rows):
        if not rows:
            return
        ws = workbook.create_sheet(sheet_name[:31])
        headers = list(rows[0].keys())
        ws.append(headers)
        for row in rows:
            ws.append([row.get(header) for header in headers])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="075E63")
            cell.alignment = Alignment(horizontal="center")
        for column_cells in ws.columns:
            width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 42)
            ws.column_dimensions[get_column_letter(column_cells[0].column)].width = max(width, 12)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    add_table("Plant Performance", [
        {"Plant": label, "Total Production": production["plant"]["total"][i],
         "OK Production": production["plant"]["ok"][i], "Rejected": production["plant"]["rejected"][i],
         "Target": production["plant"]["target"][i], "OEE %": production["plant"]["oee"][i]}
        for i, label in enumerate(production["plant"]["labels"])
    ])
    add_table("Machine Performance", production.get("machines", []))
    add_table("Quality Defects", production.get("defects", []))
    add_table("Downtime", production.get("downtime", []))
    add_table("Operator Efficiency", production.get("operators", []))
    add_table("Incidents", production.get("incidents", []))
    for cell in summary[1]:
        cell.font = Font(size=16, bold=True, color="075E63")
    for cell in summary[6]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="075E63")
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 28
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer, as_attachment=True, download_name="Production_Analytics_Summary.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/health", methods=["GET"])
def health():
    return jsonify({"ok": True, "service": "MIS Intelligence Studio"})


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5050"))
    app.run(host="0.0.0.0", port=port, debug=os.environ.get("FLASK_DEBUG") == "1", use_reloader=False)
