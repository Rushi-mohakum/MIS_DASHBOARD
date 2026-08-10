"""
parser.py — turns an uploaded file (xlsx/xls/csv/pdf/docx) into a normalized
DataFrame, then computes every number the dashboard needs.

Design goal: the person uploading data shouldn't have to match an exact
schema. Common column-name variants are auto-mapped. Missing columns get
sane defaults (and a warning) instead of crashing.
"""
import pandas as pd

# canonical_column -> list of accepted header spellings (normalized: lowercase, spaces/dashes -> underscore)
COLUMN_ALIASES = {
    "date":          ["date", "entry_date", "month", "order_date", "transaction_date", "period"],
    "product":       ["product", "item", "item_name", "sku", "product_name"],
    "category":      ["category", "product_category", "type", "segment", "department"],
    "units_sold":    ["units_sold", "qty", "quantity", "units", "qty_sold", "sold_qty"],
    "unit_price":    ["unit_price", "price", "rate", "selling_price", "mrp"],
    "net_sales":     ["net_sales", "sales", "revenue", "total_sales", "amount", "gross_sales"],
    "unit_cost":     ["unit_cost", "cost", "cost_price", "cogs_per_unit", "purchase_price"],
    "cogs":          ["cogs", "cost_of_goods_sold", "total_cost"],
    "opex":          ["opex", "operating_expense", "operating_expenses", "expenses"],
    "stock_qty":     ["stock_qty", "stock", "inventory", "current_stock", "closing_stock", "on_hand", "stock_on_hand"],
    "reorder_level": ["reorder_level", "reorder_point", "min_stock", "safety_stock", "reorder_qty"],
}


def _norm(s):
    return str(s).strip().lower().replace(" ", "_").replace("-", "_")


def normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    reverse = {}
    for canonical, aliases in COLUMN_ALIASES.items():
        for a in aliases:
            reverse[_norm(a)] = canonical
    rename_map = {c: reverse[_norm(c)] for c in df.columns if _norm(c) in reverse}
    df = df.rename(columns=rename_map)
    df = df.loc[:, ~df.columns.duplicated()]
    return df


def extract_pdf_table(path):
    import pdfplumber
    best = None
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for table in (page.extract_tables() or []):
                if not table or len(table) < 2:
                    continue
                if best is None or len(table) > len(best):
                    best = table
    if not best:
        return pd.DataFrame()
    header, *rows = best
    header = [h if h else f"col_{i}" for i, h in enumerate(header)]
    return pd.DataFrame(rows, columns=header)


def extract_docx_table(path):
    from docx import Document
    doc = Document(path)
    best = None
    for table in doc.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        if best is None or len(rows) > len(best):
            best = rows
    if not best:
        return pd.DataFrame()
    header, *rows = best
    return pd.DataFrame(rows, columns=header)


def load_dataframe(path: str, ext: str):
    """Read an uploaded file and return a normalized source DataFrame plus warnings.

    Business defaults are added later by ``prepare_business_dataframe`` so the
    data-quality and cleaning screens can still inspect the original missing
    values and data types.
    """
    warnings = []

    if ext in ("xlsx", "xls"):
        sheets = pd.read_excel(path, sheet_name=None)
        usable = {name: frame for name, frame in sheets.items() if frame is not None and not frame.empty}
        if not usable:
            return None, ["No usable worksheet was found in the workbook."]
        sheet_name, raw = max(usable.items(), key=lambda item: item[1].shape[0] * max(item[1].shape[1], 1))
        if len(usable) > 1:
            warnings.append(
                f"Workbook contains {len(usable)} non-empty sheets. Using the largest sheet: '{sheet_name}'."
            )
    elif ext == "csv":
        raw = pd.read_csv(path)
    elif ext == "pdf":
        raw = extract_pdf_table(path)
        if raw.empty:
            warnings.append("No table detected in the PDF. Use a PDF containing a real table, not a scanned image.")
    elif ext == "docx":
        raw = extract_docx_table(path)
        if raw.empty:
            warnings.append("No table detected in the Word document. Put the data inside a Word table.")
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    if raw is None or raw.empty:
        return None, warnings

    raw = raw.dropna(how="all").dropna(axis=1, how="all")
    return normalize_columns(raw), warnings


def prepare_business_dataframe(source_df: pd.DataFrame):
    """Return a dashboard-ready copy and schema warnings.

    The source DataFrame is never mutated. This keeps cleaning reversible and
    allows quality metrics to describe the actual uploaded file.
    """
    df = normalize_columns(source_df.copy())
    warnings = []

    if "product" not in df.columns:
        df["product"] = [f"Item {i+1}" for i in range(len(df))]
        warnings.append("No product column found. Generic item labels are being used.")
    if "category" not in df.columns:
        df["category"] = "Uncategorized"
        warnings.append("No category column found. All rows are grouped as Uncategorized.")
    if "units_sold" not in df.columns:
        df["units_sold"] = 0
        warnings.append("No units or quantity column found. Units Sold is shown as 0.")

    df["units_sold"] = pd.to_numeric(df["units_sold"], errors="coerce").fillna(0)

    if "net_sales" not in df.columns:
        if "unit_price" in df.columns:
            df["net_sales"] = df["units_sold"] * pd.to_numeric(df["unit_price"], errors="coerce").fillna(0)
        else:
            df["net_sales"] = 0
            warnings.append("No sales or revenue column found. Add net_sales or unit_price for meaningful sales KPIs.")

    if "cogs" not in df.columns:
        if "unit_cost" in df.columns:
            df["cogs"] = df["units_sold"] * pd.to_numeric(df["unit_cost"], errors="coerce").fillna(0)
        else:
            df["cogs"] = 0
            warnings.append("No cost column found. Add cogs or unit_cost for meaningful profit KPIs.")

    if "opex" not in df.columns:
        df["opex"] = 0
    if "stock_qty" not in df.columns:
        df["stock_qty"] = None
        warnings.append("No stock or inventory column found. Inventory KPIs are unavailable.")
    if "reorder_level" not in df.columns:
        df["reorder_level"] = None
    if "date" not in df.columns:
        df["date"] = None

    for col in ["net_sales", "cogs", "opex"]:
        df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0)
    for col in ["stock_qty", "reorder_level"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    df["product"] = df["product"].fillna("Unknown Product").astype(str)
    df["category"] = df["category"].fillna("Uncategorized").astype(str)
    return df, warnings

def compute_dashboard_data(df: pd.DataFrame, tax_rate: float = 0.25) -> dict:
    df, schema_warnings = prepare_business_dataframe(df)
    net_sales = float(df["net_sales"].sum())
    cogs = float(df["cogs"].sum())
    gross_profit = net_sales - cogs
    opex = float(df["opex"].sum())
    ebitda = gross_profit - opex
    ebit = ebitda        # no depreciation data available from a generic upload
    pbt = ebit            # no interest/other-income data available
    net_profit = pbt * (1 - tax_rate)
    units_sold = float(df["units_sold"].sum())

    gp_margin = round(gross_profit / net_sales * 100, 1) if net_sales else 0
    ebitda_margin = round(ebitda / net_sales * 100, 1) if net_sales else 0
    np_margin = round(net_profit / net_sales * 100, 1) if net_sales else 0

    # ---- Per-product rollup ----
    grouped = df.groupby("product", dropna=False).agg(
        category=("category", "first"),
        units_sold=("units_sold", "sum"),
        net_sales=("net_sales", "sum"),
        cogs=("cogs", "sum"),
        opex=("opex", "sum"),
        stock_qty=("stock_qty", "max"),
        reorder_level=("reorder_level", "max"),
    ).reset_index()
    grouped["gross_profit"] = grouped["net_sales"] - grouped["cogs"]
    grouped["net_profit"] = (grouped["gross_profit"] - grouped["opex"]) * (1 - tax_rate)
    grouped["np_margin"] = grouped.apply(
        lambda r: round(r["net_profit"] / r["net_sales"] * 100, 1) if r["net_sales"] else 0, axis=1
    )
    grouped["unit_cost_est"] = grouped.apply(
        lambda r: (r["cogs"] / r["units_sold"]) if r["units_sold"] else 0, axis=1
    )

    def _clean_records(d):
        out = []
        for rec in d:
            out.append({k: (v.item() if hasattr(v, "item") else v) for k, v in rec.items()})
        return out

    top_products = _clean_records(
        grouped.sort_values("net_sales", ascending=False).head(10).to_dict("records")
    )

    # ---- Inventory ----
    has_stock = bool(grouped["stock_qty"].notna().any())
    stock_qty_total = float(grouped["stock_qty"].fillna(0).sum())
    inventory_value = float((grouped["stock_qty"].fillna(0) * grouped["unit_cost_est"]).sum())
    denom = units_sold + stock_qty_total
    sell_through_pct = round(units_sold / denom * 100, 1) if denom else 0

    if has_stock:
        alerts_df = grouped[grouped["stock_qty"].notna() & (grouped["stock_qty"] <= grouped["reorder_level"].fillna(0))]
        stock_alerts = _clean_records(alerts_df.to_dict("records"))
    else:
        stock_alerts = []

    # ---- Category mix ----
    by_cat = df.groupby("category", dropna=False)["net_sales"].sum().sort_values(ascending=False)
    mix_labels = [str(x) for x in by_cat.index[:6]]
    mix_values = [round(float(v), 2) for v in by_cat.values[:6]]
    if len(by_cat) > 6:
        other = round(float(by_cat.values[6:].sum()), 2)
        if other:
            mix_labels.append("Other")
            mix_values.append(other)

    # ---- Monthly trend ----
    trend = {"labels": [], "net_sales": [], "gross_profit": [], "ebitda": [], "net_profit": []}
    if df["date"].notna().any():
        d = df.copy()
        d["date"] = pd.to_datetime(d["date"], errors="coerce")
        d = d.dropna(subset=["date"])
        if not d.empty:
            d["month"] = d["date"].dt.to_period("M").astype(str)
            m = d.groupby("month").agg(
                net_sales=("net_sales", "sum"),
                cogs=("cogs", "sum"),
                opex=("opex", "sum"),
            ).sort_index()
            m["gross_profit"] = m["net_sales"] - m["cogs"]
            m["ebitda"] = m["gross_profit"] - m["opex"]
            m["net_profit"] = m["ebitda"] * (1 - tax_rate)
            trend = {
                "labels": [str(x) for x in m.index],
                "net_sales": [round(float(v), 2) for v in m["net_sales"]],
                "gross_profit": [round(float(v), 2) for v in m["gross_profit"]],
                "ebitda": [round(float(v), 2) for v in m["ebitda"]],
                "net_profit": [round(float(v), 2) for v in m["net_profit"]],
            }

    # ---- Category performance (for the advanced breakdown table) ----
    cat_perf = df.groupby("category", dropna=False).agg(
        net_sales=("net_sales", "sum"),
        cogs=("cogs", "sum"),
        units_sold=("units_sold", "sum"),
    ).reset_index()
    cat_perf["gross_profit"] = cat_perf["net_sales"] - cat_perf["cogs"]
    cat_perf["gp_margin"] = cat_perf.apply(
        lambda r: round(r["gross_profit"] / r["net_sales"] * 100, 1) if r["net_sales"] else 0, axis=1
    )
    cat_perf["share_pct"] = cat_perf.apply(
        lambda r: round(r["net_sales"] / net_sales * 100, 1) if net_sales else 0, axis=1
    )
    cat_perf = _clean_records(cat_perf.sort_values("net_sales", ascending=False).to_dict("records"))

    # ---- Month-over-month growth (needs at least 2 months of trend data) ----
    momentum = {"net_sales_growth_pct": None, "units_growth_pct": None, "has_momentum": False}
    if len(trend["labels"]) >= 2:
        prev_sales, last_sales = trend["net_sales"][-2], trend["net_sales"][-1]
        if prev_sales:
            momentum["net_sales_growth_pct"] = round((last_sales - prev_sales) / prev_sales * 100, 1)
        momentum["has_momentum"] = True
        momentum["last_month"] = trend["labels"][-1]
        momentum["prev_month"] = trend["labels"][-2]

    return {
        "kpi": {
            "net_sales": net_sales,
            "gross_profit": gross_profit,
            "gp_margin": gp_margin,
            "ebitda": ebitda,
            "ebitda_margin": ebitda_margin,
            "net_profit": net_profit,
            "np_margin": np_margin,
            "units_sold": units_sold,
            "inventory_value": inventory_value,
            "sell_through_pct": sell_through_pct,
            "stock_alerts_count": len(stock_alerts),
            "has_stock": has_stock,
            "sku_count": len(grouped),
        },
        "pnl": {
            "gross_sales": net_sales,   # no returns/discounts data in a generic upload
            "net_sales": net_sales,
            "cogs": cogs,
            "gross_profit": gross_profit,
            "opex": opex,
            "ebitda": ebitda,
            "ebit": ebit,
            "pbt": pbt,
            "net_profit": net_profit,
        },
        "trend": trend,
        "momentum": momentum,
        "mix_labels": mix_labels,
        "mix_values": mix_values,
        "top_products": top_products,
        "category_performance": cat_perf,
        "stock_alerts": stock_alerts,
        "schema_warnings": schema_warnings,
    }


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

def build_excel_report(df: pd.DataFrame, data: dict, tax_rate: float) -> "BytesIO":
    from io import BytesIO
    from datetime import datetime
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    HEADER_FILL = PatternFill(start_color="1F4E5F", end_color="1F4E5F", fill_type="solid")
    HEADER_FONT = Font(color="FFFFFF", bold=True)
    TITLE_FONT = Font(size=16, bold=True, color="1F4E5F")
    SUBTITLE_FONT = Font(size=10, italic=True, color="666666")
    LABEL_FONT = Font(bold=True)
    THIN_BORDER = Border(*(Side(style="thin", color="DDDDDD"),) * 4)
    CENTER = Alignment(horizontal="center", vertical="center")

    def style_header(ws, row, ncols):
        for col in range(1, ncols + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = CENTER
            cell.border = THIN_BORDER

    def autosize(ws):
        for col in ws.columns:
            length = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(length + 3, 40)

    def write_table(ws, start_row, headers, rows):
        for i, h in enumerate(headers, start=1):
            ws.cell(row=start_row, column=i, value=h)
        style_header(ws, start_row, len(headers))
        for r_idx, row in enumerate(rows, start=start_row + 1):
            for c_idx, val in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.border = THIN_BORDER
        return start_row + len(rows) + 1

    wb = Workbook()
    kpi, pnl = data["kpi"], data["pnl"]

    # ---- Summary sheet ----
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "MIS Dashboard — Summary"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}  |  Assumed tax rate {round(tax_rate*100)}%"
    ws["A2"].font = SUBTITLE_FONT
    ws.merge_cells("A1:D1")
    ws.merge_cells("A2:D2")

    rows = [
        ("Net Sales", round(kpi["net_sales"], 2)),
        ("Gross Profit", round(kpi["gross_profit"], 2)),
        ("GP Margin %", kpi["gp_margin"]),
        ("EBITDA", round(kpi["ebitda"], 2)),
        ("EBITDA Margin %", kpi["ebitda_margin"]),
        ("Net Profit", round(kpi["net_profit"], 2)),
        ("NP Margin %", kpi["np_margin"]),
        ("Units Sold", round(kpi["units_sold"], 2)),
        ("Inventory Value", round(kpi["inventory_value"], 2)),
        ("Sell-Through Rate %", kpi["sell_through_pct"]),
        ("Stock Alerts", kpi["stock_alerts_count"]),
        ("Products Tracked", kpi["sku_count"]),
    ]
    row = 4
    ws.cell(row=row, column=1, value="Metric").font = LABEL_FONT
    ws.cell(row=row, column=2, value="Value").font = LABEL_FONT
    style_header(ws, row, 2)
    for i, (label, val) in enumerate(rows, start=row + 1):
        ws.cell(row=i, column=1, value=label).border = THIN_BORDER
        ws.cell(row=i, column=2, value=val).border = THIN_BORDER
    autosize(ws)

    # ---- P&L sheet ----
    ws2 = wb.create_sheet("P&L Snapshot")
    ws2["A1"] = "Profit & Loss Snapshot"
    ws2["A1"].font = TITLE_FONT
    ws2["A2"] = "Gross Sales assumed = Net Sales; EBIT/PBT assumed = EBITDA (generic upload has no returns/depreciation/interest data)."
    ws2["A2"].font = SUBTITLE_FONT
    pnl_rows = [
        ("Gross Sales", round(pnl["gross_sales"], 2)),
        ("Net Sales", round(pnl["net_sales"], 2)),
        ("COGS", round(pnl["cogs"], 2)),
        ("Gross Profit", round(pnl["gross_profit"], 2)),
        ("OpEx", round(pnl["opex"], 2)),
        ("EBITDA", round(pnl["ebitda"], 2)),
        ("EBIT", round(pnl["ebit"], 2)),
        ("PBT", round(pnl["pbt"], 2)),
        ("Net Profit", round(pnl["net_profit"], 2)),
    ]
    write_table(ws2, 4, ["Line Item", "Amount"], pnl_rows)
    autosize(ws2)

    # ---- Top Products sheet ----
    ws3 = wb.create_sheet("Top Products")
    headers = ["Product", "Category", "Units Sold", "Net Sales", "Gross Profit", "Net Profit", "NP Margin %"]
    prod_rows = [
        [p["product"], p["category"], round(p["units_sold"], 2), round(p["net_sales"], 2),
         round(p["gross_profit"], 2), round(p["net_profit"], 2), p["np_margin"]]
        for p in data["top_products"]
    ]
    if not prod_rows:
        prod_rows = [["No product data", "", "", "", "", "", ""]]
    write_table(ws3, 1, headers, prod_rows)
    autosize(ws3)

    # ---- Category performance sheet ----
    ws4 = wb.create_sheet("Category Performance")
    cat_headers = ["Category", "Net Sales", "Gross Profit", "GP Margin %", "Share of Sales %", "Units Sold"]
    cat_rows = [
        [c["category"], round(c["net_sales"], 2), round(c["gross_profit"], 2), c["gp_margin"], c["share_pct"], round(c["units_sold"], 2)]
        for c in data["category_performance"]
    ]
    if not cat_rows:
        cat_rows = [["No category data", "", "", "", "", ""]]
    write_table(ws4, 1, cat_headers, cat_rows)
    autosize(ws4)

    # ---- Monthly trend sheet ----
    ws5 = wb.create_sheet("Monthly Trend")
    if data["trend"]["labels"]:
        t_headers = ["Month", "Net Sales", "Gross Profit", "EBITDA", "Net Profit"]
        t_rows = list(zip(
            data["trend"]["labels"], data["trend"]["net_sales"],
            data["trend"]["gross_profit"], data["trend"]["ebitda"], data["trend"]["net_profit"],
        ))
        write_table(ws5, 1, t_headers, [list(r) for r in t_rows])
    else:
        ws5["A1"] = "No date column found in the uploaded file — no monthly trend available."
    autosize(ws5)

    # ---- Inventory alerts sheet ----
    ws6 = wb.create_sheet("Inventory Alerts")
    if data["stock_alerts"]:
        a_headers = ["Product", "Category", "Stock Qty", "Reorder Level"]
        a_rows = [[a["product"], a["category"], a["stock_qty"], a.get("reorder_level")] for a in data["stock_alerts"]]
        write_table(ws6, 1, a_headers, a_rows)
    else:
        ws6["A1"] = "No products at or below reorder level."
    autosize(ws6)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# AI insights prompt
# ---------------------------------------------------------------------------

def build_insights_prompt(data: dict, tax_rate: float) -> str:
    """Compact, numbers-only summary handed to the model — keeps the prompt small
    and stops the model from inventing figures that weren't in the data."""
    import json

    kpi, pnl = data["kpi"], data["pnl"]
    summary = {
        "kpis": {
            "net_sales": round(kpi["net_sales"], 2),
            "gross_profit": round(kpi["gross_profit"], 2),
            "gp_margin_pct": kpi["gp_margin"],
            "ebitda": round(kpi["ebitda"], 2),
            "ebitda_margin_pct": kpi["ebitda_margin"],
            "net_profit": round(kpi["net_profit"], 2),
            "np_margin_pct": kpi["np_margin"],
            "units_sold": round(kpi["units_sold"], 2),
            "inventory_value": round(kpi["inventory_value"], 2) if kpi["has_stock"] else None,
            "sell_through_pct": kpi["sell_through_pct"] if kpi["has_stock"] else None,
            "stock_alerts_count": kpi["stock_alerts_count"],
            "assumed_tax_rate_pct": round(tax_rate * 100),
        },
        "monthly_trend": {
            "months": data["trend"]["labels"],
            "net_sales": data["trend"]["net_sales"],
        } if data["trend"]["labels"] else None,
        "category_performance": [
            {"category": c["category"], "net_sales": round(c["net_sales"], 2),
             "gp_margin_pct": c["gp_margin"], "share_of_sales_pct": c["share_pct"]}
            for c in data["category_performance"]
        ],
        "top_products": [
            {"product": p["product"], "category": p["category"], "net_sales": round(p["net_sales"], 2),
             "np_margin_pct": p["np_margin"]}
            for p in data["top_products"][:8]
        ],
        "stock_alerts": [
            {"product": s["product"], "category": s["category"], "stock_qty": s["stock_qty"]}
            for s in data["stock_alerts"][:10]
        ],
    }

    return f"""You are a business analyst reviewing an MIS dashboard for a manufacturing/sales business. \
Base every claim strictly on the JSON data below — never invent numbers, competitors, or market names that \
aren't implied by the categories given. If the data is too thin for a section, say so briefly instead of \
guessing.

DATA:
{json.dumps(summary, indent=2)}

Respond in exactly these four short sections, using markdown headers (####) and concise bullet points \
(3-5 bullets per section, no fluff):

#### KPIs to watch closely
Which of the metrics above most need monitoring right now, and why (e.g. thin margins, declining trend, \
concentration risk in one category/product).

#### Where output/volume could improve
Practical suggestions to grow units sold or sell-through, grounded in what the category/product/stock data \
actually shows (e.g. products with strong margin but low volume, categories with high stock and low sell-through).

#### Where to focus investment
Based only on the category performance figures given (margin % and share of sales), which category or \
categories look most worth doubling down on, and which look weakest. Do not invent geographic markets that \
aren't in the data — talk about product/category segments instead unless the data includes real market/region info.

#### Risks & data gaps
Anything in the numbers that looks risky (e.g. stock alerts, thin margins, only one month of data) and what \
additional data would make this analysis more reliable.
"""
